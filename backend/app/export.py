"""Export researched leads and their emails to DOCX or PDF."""

import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

WINDOWS_FONT = Path("C:/Windows/Fonts/arial.ttf")


def _sections(lead: dict[str, Any]) -> list[tuple[str, str, str]]:
    """(label, subject, body) for each email in the sequence."""
    return [
        ("Email 1", lead.get("subject") or "", lead.get("body") or ""),
        ("Follow-up 1", lead.get("follow_up_1_subject") or "", lead.get("follow_up_1_body") or ""),
        ("Follow-up 2", lead.get("follow_up_2_subject") or "", lead.get("follow_up_2_body") or ""),
    ]


def _heading(lead: dict[str, Any]) -> str:
    name = f"{lead.get('first_name', '')} {lead.get('last_name', '')}".strip()
    return f"{name} - {lead.get('company') or ''}".strip(" -")


def _meta(lead: dict[str, Any]) -> str:
    parts = [lead.get("title"), lead.get("email"), lead.get("website"), lead.get("city")]
    return " | ".join(p for p in parts if p)


def _strip_md(text: str) -> str:
    return re.sub(r"[*#`]+", "", text or "")


def to_docx(leads: list[dict[str, Any]], include_research: bool) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Cold email drafts", 0)
    for i, lead in enumerate(leads):
        if i:
            doc.add_page_break()
        doc.add_heading(_heading(lead), 1)
        doc.add_paragraph().add_run(_meta(lead)).italic = True
        if lead.get("personalization_notes"):
            p = doc.add_paragraph()
            p.add_run("Personalization: ").bold = True
            p.add_run(lead["personalization_notes"])
        for label, subject, body in _sections(lead):
            doc.add_heading(label, 2)
            p = doc.add_paragraph()
            p.add_run("Subject: ").bold = True
            p.add_run(subject)
            doc.add_paragraph(body)
        if include_research and lead.get("research"):
            doc.add_heading("Research brief", 2)
            for line in lead["research"].splitlines():
                if line.startswith("## "):
                    doc.add_heading(line[3:], 3)
                elif line.strip():
                    doc.add_paragraph(_strip_md(line.lstrip("- ")), style="List Bullet" if line.lstrip().startswith("- ") else None)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(leads: list[dict[str, Any]], include_research: bool) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=15)
    if WINDOWS_FONT.exists():
        pdf.add_font("Body", "", str(WINDOWS_FONT))
        pdf.add_font("Body", "B", str(WINDOWS_FONT.with_name("arialbd.ttf")))
        pdf.add_font("Body", "I", str(WINDOWS_FONT.with_name("ariali.ttf")))
        font, clean = "Body", (lambda s: s or "")
    else:  # Core fonts are latin-1 only.
        font = "Helvetica"
        clean = lambda s: (s or "").encode("latin-1", "replace").decode("latin-1")  # noqa: E731

    def line(text: str, size: int = 11, style: str = "", gap: float = 2) -> None:
        pdf.set_font(font, style, size)
        pdf.multi_cell(0, size * 0.5, clean(text))
        pdf.ln(gap)

    for lead in leads:
        pdf.add_page()
        line(_heading(lead), 16, "B")
        line(_meta(lead), 9, "I", 4)
        if lead.get("personalization_notes"):
            line("Personalization: " + lead["personalization_notes"], 9, "I", 5)
        for label, subject, body in _sections(lead):
            line(label, 13, "B", 1)
            line("Subject: " + subject, 11, "B")
            line(body, 11, "", 6)
        if include_research and lead.get("research"):
            line("Research brief", 13, "B")
            for raw in lead["research"].splitlines():
                if raw.startswith("## "):
                    line(raw[3:], 11, "B", 1)
                elif raw.strip():
                    line(_strip_md(raw), 9, "", 1)
    return bytes(pdf.output())
